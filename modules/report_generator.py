import pandas as pd
import openpyxl

def generate_excel(data: list, report_path: str):
    """Generates a multi-sheet Excel report with inspection findings, summary, and category stats."""
    findings = pd.DataFrame(data)

    if findings.empty:
        summary = pd.DataFrame([["Total Findings", 0]], columns=["Metric", "Value"])
        category = pd.DataFrame(columns=["Issue Type", "Count"])
    else:
        summary = pd.DataFrame(
            [
                ["Total Findings", len(findings)],
                ["High Severity", len(findings[findings["Severity"] == "High"])],
                ["Medium Severity", len(findings[findings["Severity"] == "Medium"])],
                ["Low Severity", len(findings[findings["Severity"] == "Low"])],
            ],
            columns=["Metric", "Value"]
        )

        category = findings.groupby("Issue Type").size().reset_index()
        category.columns = ["Issue Type", "Count"]

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        findings.to_excel(writer, sheet_name="Inspection Findings", index=False)
        summary.to_excel(writer, sheet_name="Executive Summary", index=False)
        category.to_excel(writer, sheet_name="Categories", index=False)

    # Post-process formatting: adjust column widths automatically
    wb = openpyxl.load_workbook(report_path)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
    wb.save(report_path)
